"""Smaller sources: NS IFF timetable station attributes and ProRail's NDOV files.

All from NDOV Loket (CC0): https://data.ndovloket.nl/
Outputs (data/raw/derived/):
  iff_stations.parquet          NS station codes, country, TGST/RAST attributes
  prorail_2020_stepfree.parquet ProRail step-free map input (2019/2020 snapshot)
  prorail_boarding.parquet      ProRail per-platform boarding-height status (Q2 2026)
  prorail_assets.parquet        lifts and ramps per station in ProRail's asset register (June 2026)
"""

from __future__ import annotations

import zipfile
from collections import defaultdict
from urllib.parse import quote

import openpyxl
import polars as pl

from fetch import RAW, fetch

DERIVED = RAW / "derived"
PRORAIL = "https://data.ndovloket.nl/prorail/"
STEPFREE_2020 = "Drempelvrije toegankelijkheid stations - input kaart.xlsx"
BOARDING = "Analyseblad landelijk en Toegankelijkheid Samenvatting Q2 2026.xlsx"
LIFTS = "Liften juni 2026.csv"
RAMPS = "Hellingbanen juni 2026.csv"


def iff_stations() -> pl.DataFrame:
    z = zipfile.ZipFile(fetch("https://data.ndovloket.nl/ns/ns-latest.zip", "ns-latest.zip"))
    attrs: dict[str, set[str]] = defaultdict(set)
    cur = None
    for line in z.read("attributesonstation.dat").decode("latin1").splitlines():
        if line.startswith("#"):
            cur = line[1:].strip()
        elif line.startswith("-") and cur:
            attrs[cur].add(line[1:].strip())
    rows = []
    for line in z.read("stations.dat").decode("latin1").splitlines()[1:]:
        f = [x.strip() for x in line.split(",")]
        code = f[1]
        rows.append(
            {
                "ns_code": code.upper(),
                "iff_name": f[9],
                "iff_country": f[4],
                "iff_transfer_station": f[0] == "1",
                "iff_transfer_minutes": int(f[2]),
                "ns_tgst": "TGST" in attrs.get(code, ()),
                "ns_rast": "RAST" in attrs.get(code, ()),
            }
        )
    header = z.read("stations.dat").decode("latin1").splitlines()[0]
    return pl.DataFrame(rows).with_columns(pl.lit(header.split(",")[1]).alias("iff_valid_from"))


def prorail_2020() -> pl.DataFrame:
    """ProRail's input for its step-free station map.

    Two sheets. 'Uitrollijst (ex AVG)' is the newer one (edits up to 2020-07-21),
    but its first column ('Station') was sorted separately from the rest:
    only StationCorrectie / Afkorting line up with the status, date and note
    columns, and the last few stations fell off. 'Uitrollijst' is consistent
    but older (2019). Use the newer sheet, fall back to the older one.
    """
    p = fetch(PRORAIL + quote(STEPFREE_2020), "prorail/" + STEPFREE_2020)
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    reach = {"Ja": "yes", "Nee": "no"}
    rows: dict[str, dict] = {}
    for r in list(wb["Uitrollijst (ex AVG)"].iter_rows(values_only=True))[1:]:
        if not r[0] or not r[2]:
            continue
        rows[str(r[2]).strip().upper()] = {
            "pr2020_name": str(r[1]).strip(),
            "pr2020_reachable": reach.get(r[3], "unknown"),
            "pr2020_category": r[4],
            "pr2020_date": r[6].date().isoformat() if r[6] else None,
            "pr2020_note": r[7],
            "pr2020_sheet": "ex AVG",
        }
    for r in list(wb["Uitrollijst"].iter_rows(values_only=True))[1:]:
        code = str(r[2]).strip().upper() if r[2] else None
        if not code or code in rows:
            continue
        rows[code] = {
            "pr2020_name": str(r[1]).strip(),
            "pr2020_reachable": reach.get(r[4], "unknown"),
            "pr2020_category": r[5],
            "pr2020_date": None,
            "pr2020_note": None,
            "pr2020_sheet": "Uitrollijst",
        }
    return pl.DataFrame([{"ns_code": k, **v} for k, v in rows.items()], infer_schema_length=None)


def prorail_boarding() -> pl.DataFrame:
    p = fetch(PRORAIL + quote(BOARDING), "prorail/" + BOARDING)
    ws = openpyxl.load_workbook(p, read_only=True, data_only=True).worksheets[0]
    rows = [
        {"pr_board_name": r[0], "track": str(r[1]), "status": r[4], "carriers": r[7]}
        for r in list(ws.iter_rows(values_only=True))[2:]
        if r[0]
    ]
    df = pl.DataFrame(rows, infer_schema_length=None)
    return df.group_by("pr_board_name").agg(
        pl.len().alias("pr_board_tracks"),
        (pl.col("status") == "Toegankelijk").sum().alias("pr_board_tracks_ok"),
        pl.col("status").unique().sort().str.join("|").alias("pr_board_statuses"),
    )


def prorail_register(filename: str) -> pl.DataFrame:
    """ProRail asset register export (SAP), semicolon CSV with two header rows."""
    p = fetch(PRORAIL + quote(filename), "prorail/" + filename)
    return pl.read_csv(
        p, separator=";", skip_rows=1, infer_schema=False, encoding="utf8-lossy", truncate_ragged_lines=True
    ).with_columns(
        pl.col("INR_VERKORTING").str.strip_chars().str.to_uppercase().alias("ns_code"),
        pl.col("STTXU").str.strip_chars().alias("user_status"),
    )


def prorail_assets() -> pl.DataFrame:
    lifts = prorail_register(LIFTS).group_by("ns_code").agg(
        pl.len().alias("reg_lifts"),
        pl.col("user_status").unique().sort().str.join("|").alias("reg_lift_status"),
    )
    ramps = prorail_register(RAMPS).group_by("ns_code").agg(pl.len().alias("reg_ramps"))
    return lifts.join(ramps, on="ns_code", how="full", coalesce=True)


def main() -> None:
    DERIVED.mkdir(parents=True, exist_ok=True)
    iff_stations().write_parquet(DERIVED / "iff_stations.parquet")
    prorail_2020().write_parquet(DERIVED / "prorail_2020_stepfree.parquet")
    prorail_boarding().write_parquet(DERIVED / "prorail_boarding.parquet")
    prorail_assets().write_parquet(DERIVED / "prorail_assets.parquet")
    print("misc sources written")


if __name__ == "__main__":
    main()
